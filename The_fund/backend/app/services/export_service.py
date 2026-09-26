# app/services/export_service.py
"""
توليد تقارير PDF وExcel بعربية سليمة.

التغييرات عن النسخة السابقة (ملاحظتا الفحص 31 و32):
  - لا تنزيل أي شيء من الإنترنت عند الاستيراد أو عند إعداد التقرير. الخطوط
    تُحمَّل من مجلد المشروع app/static/fonts فقط (ضع Amiri-Regular.ttf و
    Amiri-Bold.ttf هناك أثناء النشر).
  - في الإنتاج، غياب الخط يفشل فوراً برسالة إعداد واضحة بدل إنتاج PDF
    بعربية غير مقروءة. في التطوير يُستخدم Helvetica مع تحذير.
"""
import io
import os
import logging
from datetime import datetime

from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from app.core.config import settings

logger = logging.getLogger("export_service")

BASE_DIR = os.path.dirname(os.path.dirname(__file__))
FONT_DIR = os.path.join(BASE_DIR, "static", "fonts")
FONT_PATH = os.path.join(FONT_DIR, "Amiri-Regular.ttf")
FONT_BOLD_PATH = os.path.join(FONT_DIR, "Amiri-Bold.ttf")

try:
    from arabic_reshaper import ArabicReshaper
    _reshaper = ArabicReshaper(configuration={"delete_harakat": False, "support_ligatures": True})
    from bidi.algorithm import get_display
    _SHAPING_AVAILABLE = True
except ImportError:
    _reshaper = None
    get_display = None
    _SHAPING_AVAILABLE = False
    logger.warning("arabic_reshaper/python-bidi غير مثبتة - النص العربي في PDF لن يُشكَّل بشكل صحيح.")

_fonts_ready = False
ARABIC_FONT = "Helvetica"
ARABIC_FONT_BOLD = "Helvetica-Bold"


def _ensure_font_registered():
    """يسجّل الخطوط العربية من مجلد المشروع (بدون أي وصول للشبكة)."""
    global _fonts_ready, ARABIC_FONT, ARABIC_FONT_BOLD
    if _fonts_ready:
        return True

    if not os.path.exists(FONT_PATH):
        message = (
            f"الخط العربي غير موجود: {FONT_PATH}. ضع Amiri-Regular.ttf و Amiri-Bold.ttf في {FONT_DIR} "
            "أثناء النشر (حمّلهما مرة واحدة من https://fonts.google.com/specimen/Amiri واحفظهما في المشروع)."
        )
        if settings.is_production:
            raise RuntimeError(message)
        logger.warning(message + " (وضع التطوير: سيُستخدم Helvetica وسيظهر النص العربي غير سليم)")
        return False

    try:
        pdfmetrics.registerFont(TTFont("Arabic", FONT_PATH))
        if os.path.exists(FONT_BOLD_PATH):
            pdfmetrics.registerFont(TTFont("Arabic-Bold", FONT_BOLD_PATH))
            ARABIC_FONT, ARABIC_FONT_BOLD = "Arabic", "Arabic-Bold"
        else:
            ARABIC_FONT, ARABIC_FONT_BOLD = "Arabic", "Arabic"
        _fonts_ready = True
        return True
    except Exception as e:
        logger.warning("تعذر تسجيل الخط العربي مع reportlab: %s", e)
        if settings.is_production:
            raise RuntimeError(f"تعذر تحميل الخط العربي: {e}")
        return False


def shape_arabic(text):
    if not text:
        return text
    text = str(text)
    if not _SHAPING_AVAILABLE:
        return text
    try:
        reshaped = _reshaper.reshape(text)
        return get_display(reshaped)
    except Exception:
        return text


def _pdf_styles():
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="ArabicTitle", fontName=ARABIC_FONT_BOLD, fontSize=16,
                              alignment=1, spaceAfter=10))
    styles.add(ParagraphStyle(name="ArabicNormal", fontName=ARABIC_FONT, fontSize=10, alignment=1))
    return styles


def generate_pdf_report(title, headers, rows, fund_info=None):
    _ensure_font_registered()

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, topMargin=1.5 * cm, bottomMargin=1.5 * cm)
    styles = _pdf_styles()
    elements = []

    if fund_info:
        elements.append(Paragraph(shape_arabic(fund_info.get("name", "")), styles["ArabicTitle"]))
        contact = " | ".join(filter(None, [fund_info.get("phone"), fund_info.get("email")]))
        if contact:
            elements.append(Paragraph(shape_arabic(contact), styles["ArabicNormal"]))
        elements.append(Spacer(1, 12))

    elements.append(Paragraph(shape_arabic(title), styles["ArabicTitle"]))
    elements.append(Paragraph(datetime.now().strftime("%Y-%m-%d %H:%M"), styles["ArabicNormal"]))
    elements.append(Spacer(1, 12))

    shaped_headers = [shape_arabic(h) for h in headers]
    shaped_rows = [[shape_arabic(cell) for cell in row] for row in rows]
    table_data = [shaped_headers] + shaped_rows

    table = Table(table_data, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1B5E20")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), ARABIC_FONT_BOLD),
        ("FONTNAME", (0, 1), (-1, -1), ARABIC_FONT),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CCCCCC")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5FAF6")]),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    elements.append(table)

    doc.build(elements)
    buffer.seek(0)
    return buffer.read()


def generate_excel_report(title, headers, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.sheet_view.rightToLeft = True

    header_fill = PatternFill(start_color="1B5E20", end_color="1B5E20", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=11)

    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row_idx, row in enumerate(rows, start=2):
        for col_idx, value in enumerate(row, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = Alignment(horizontal="center")
            if row_idx % 2 == 0:
                cell.fill = PatternFill(start_color="F5FAF6", end_color="F5FAF6", fill_type="solid")

    for col_idx in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 20

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.read()
